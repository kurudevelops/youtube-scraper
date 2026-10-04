"""
YouTube collector for MATATAG / K-12 opinion scraping -> Excel output.

Setup:
    pip install google-api-python-client openpyxl python-dotenv
    .env:
        YOUTUBE_API_KEY=your_key

Run:
    python youtube_collector.py --out data.xlsx
    python youtube_collector.py --max-videos 15 --max-comments 300 --only-tagged

Rerunning appends only new videos/comments (deduped by ID). Close the
Excel file before running, or the save will fail.
"""
import argparse
import html
import os
import re

from dotenv import load_dotenv
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

load_dotenv()

# Opinion-oriented queries: reactions, debates, pros/cons, news with public
# discussion. Tutorials and explainers attract "thank you po" comments, not opinions.
SEARCH_QUERIES = [
    "MATATAG curriculum teachers react",
    "MATATAG curriculum pros and cons",
    "MATATAG curriculum problema teachers",
    "MATATAG curriculum workload teachers review",
    "MATATAG kurikulum reaksyon ng guro",
    "epekto ng MATATAG sa mga estudyante",
    "DepEd MATATAG news",
    "K to 12 curriculum failed Philippines",
    "K-12 program pros and cons Philippines",
    "senior high school removal Philippines opinion",
    "K to 12 kurikulum opinyon",
]

MATATAG_RE = re.compile(r"matatag|kto10|k\s?to\s?10", re.I)
K12_RE = re.compile(r"k[\s-]?12|k\s?to\s?12|senior\s?high|\bshs\b", re.I)
# A video must mention one of these in its title/description to be kept.
RELEVANT_RE = re.compile(
    r"matatag|kto10|k[\s-]?12|k\s?to\s?1[02]|senior\s?high|\bshs\b|deped|curriculum|kurikulum",
    re.I)
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # openpyxl rejects these

VIDEO_COLS = ["video_id", "title", "channel_title", "published_at", "query", "comment_count"]
COMMENT_COLS = ["comment_id", "video_id", "parent_id", "text",
                "like_count", "published_at", "topic_tag"]
# Author names are deliberately not stored (privacy, if results get published).

MIN_COMMENT_CHARS = 25  # drops "thank you po", "first", emoji-only comments


def clean(v):
    """Excel cells: strip illegal control chars, cap at the 32,767-char limit."""
    return ILLEGAL_RE.sub("", v)[:32000] if isinstance(v, str) else v


def tag_topic(text):
    m, k = bool(MATATAG_RE.search(text)), bool(K12_RE.search(text))
    return "both" if m and k else "matatag" if m else "k12" if k else "none"


def open_workbook(path):
    """Load existing file (and the IDs already in it) or start a new one."""
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    sheets = {}
    for name, cols in (("videos", VIDEO_COLS), ("comments", COMMENT_COLS)):
        if name in wb.sheetnames:
            ws = wb[name]
        else:
            ws = wb.create_sheet(name)
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True)
            ws.freeze_panes = "A2"
        sheets[name] = ws
    seen_videos = {r[0] for r in sheets["videos"].iter_rows(min_row=2, values_only=True)}
    seen_comments = {r[0] for r in sheets["comments"].iter_rows(min_row=2, values_only=True)}
    return wb, sheets, seen_videos, seen_comments


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def search_videos(yt, query, max_videos):
    """search.list costs 100 quota units per call; keep queries few."""
    found, token = [], None
    while len(found) < max_videos:
        resp = yt.search().list(
            q=query, part="snippet", type="video",
            maxResults=min(50, max_videos - len(found)),
            regionCode="PH", relevanceLanguage="en", pageToken=token,
        ).execute()
        for item in resp["items"]:
            s = item["snippet"]
            found.append({
                "id": item["id"]["videoId"],
                "title": html.unescape(s["title"]),
                "channel": s["channelTitle"],
                "published": s["publishedAt"],
                "desc": s.get("description", ""),
                "query": query,
            })
        token = resp.get("nextPageToken")
        if not token:
            break
    return found


def comment_counts(yt, ids):
    """videos.list costs 1 unit per call (50 IDs). Missing count = comments off."""
    counts = {}
    for i in range(0, len(ids), 50):
        resp = yt.videos().list(part="statistics", id=",".join(ids[i:i + 50])).execute()
        for item in resp["items"]:
            counts[item["id"]] = int(item["statistics"].get("commentCount", 0))
    return counts


def parse_comment(c, video_id, parent_id):
    s = c["snippet"]
    text = s["textDisplay"]
    return (c["id"], video_id, parent_id, text, s.get("likeCount", 0),
            s["publishedAt"], tag_topic(text))


def fetch_comments(yt, video_id, max_comments):
    """commentThreads.list costs 1 unit per call (100 threads max)."""
    rows, token = [], None
    while len(rows) < max_comments:
        try:
            resp = yt.commentThreads().list(
                part="snippet,replies", videoId=video_id, maxResults=100,
                order="relevance", textFormat="plainText", pageToken=token,
            ).execute()
        except HttpError as e:
            if e.resp.status in (403, 404):  # comments disabled / video gone
                print(f"  skip {video_id}: {e.resp.status}")
                return rows
            raise
        for thread in resp["items"]:
            top = thread["snippet"]["topLevelComment"]
            rows.append(parse_comment(top, video_id, None))
            for reply in thread.get("replies", {}).get("comments", []):
                rows.append(parse_comment(reply, video_id, top["id"]))
        token = resp.get("nextPageToken")
        if not token:
            break
    return rows[:max_comments]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--max-videos", type=int, default=15, help="search results per query")
    ap.add_argument("--max-comments", type=int, default=300, help="per video")
    ap.add_argument("--min-video-comments", type=int, default=30,
                    help="skip videos with fewer total comments than this")
    ap.add_argument("--only-tagged", action="store_true",
                    help="keep only comments that name MATATAG / K-12 explicitly")
    ap.add_argument("--out", default="matatag_youtube.xlsx")
    args = ap.parse_args()

    yt = build("youtube", "v3", developerKey=os.environ["YOUTUBE_API_KEY"])
    wb, sheets, seen_videos, seen_comments = open_workbook(args.out)

    for q in SEARCH_QUERIES:
        print(f"[search] {q}")
        candidates = [v for v in search_videos(yt, q, args.max_videos)
                      if v["id"] not in seen_videos
                      and RELEVANT_RE.search(v["title"] + " " + v["desc"])]
        if not candidates:
            continue
        counts = comment_counts(yt, [v["id"] for v in candidates])
        for v in candidates:
            n = counts.get(v["id"], 0)
            if n < args.min_video_comments:
                print(f"  drop {v['id']} ({n} comments): {v['title'][:50]}")
                continue
            seen_videos.add(v["id"])
            sheets["videos"].append([clean(x) for x in (
                v["id"], v["title"], v["channel"], v["published"], v["query"], n)])
            new = 0
            for row in fetch_comments(yt, v["id"], args.max_comments):
                if row[0] in seen_comments:
                    continue
                if len(row[3]) < MIN_COMMENT_CHARS:
                    continue
                if args.only_tagged and row[6] == "none":
                    continue
                seen_comments.add(row[0])
                sheets["comments"].append([clean(x) for x in row])
                new += 1
            wb.save(args.out)  # save per video so a crash/quota stop loses nothing
            print(f"  {v['id']}: {new} new comments | {v['title'][:60]}")

    set_widths(sheets["videos"], [14, 60, 28, 22, 30, 14])
    set_widths(sheets["comments"], [28, 14, 28, 90, 10, 22, 12])
    wb.save(args.out)
    print(f"Saved -> {args.out}")


if __name__ == "__main__":
    main()